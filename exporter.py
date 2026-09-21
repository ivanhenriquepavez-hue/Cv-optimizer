import io
import docx

def generate_optimized_cv_docx(original_text: str, new_bullets: list) -> io.BytesIO:
    """
    Genera un documento DOCX en memoria que fusiona el contenido del CV original 
    con los nuevos puntos clave (STAR) generados, para que el usuario pueda descargarlo.
    """
    try:
        doc = docx.Document()
        
        # Título del documento
        doc.add_heading('Currículum Vitae (Optimizado ATS)', 0)
        
        # --- SECCIÓN 1: Puntos Nuevos Generados ---
        doc.add_heading('1. Nuevas Competencias y Logros (Formato STAR)', level=1)
        p_intro = doc.add_paragraph("Las siguientes experiencias han sido estructuradas para resaltar métricas y resultados, cubriendo las brechas detectadas por el sistema ATS.")
        p_intro.style = 'Italic'
        
        for item in new_bullets:
            p = doc.add_paragraph(style='List Bullet')
            p.add_run(f"{item['competencia']}: ").bold = True
            p.add_run(item['bullet'])
            
        doc.add_paragraph() # Espacio
        doc.add_page_break() # Salto de página para separar contenido
        
        # --- SECCIÓN 2: Texto Original ---
        doc.add_heading('2. Contenido Base (CV Original)', level=1)
        
        for line in original_text.split('\n'):
            if line.strip():
                doc.add_paragraph(line.strip())
                
        # Guardamos en un buffer en RAM para no escribir a disco temporalmente,
        # esto es ideal para descargas directas en Streamlit.
        file_stream = io.BytesIO()
        doc.save(file_stream)
        file_stream.seek(0)
        
        return file_stream
        
    except Exception as e:
        raise Exception(f"Error generando el documento DOCX: {str(e)}")
